from flask import Flask, request, jsonify
import subprocess
import os
from datetime import datetime

app = Flask(__name__)

# قائمة الأوامر المسموحة (آمن)
ALLOWED_APPS = {
    "firefox": "firefox",
    "chrome": "google-chrome",
    "terminal": "x-terminal-emulator",
    "files": "thunar",
    "calculator": "gnome-calculator",
    "text editor": "gedit",
    "code": "code",
    "vlc": "vlc",
}

@app.route('/execute', methods=['POST'])
def execute():
    data = request.json
    action = data.get('action')
    target = data.get('target', '')
    
    # ============ أوامر التطبيقات ============
    if action == 'open_app':
        app_cmd = ALLOWED_APPS.get(target.lower())
        if not app_cmd:
            return jsonify({
                "success": False,
                "message": f"التطبيق '{target}' غير مسموح"
            })
        
        try:
            subprocess.Popen([app_cmd])
            return jsonify({
                "success": True,
                "message": f"✅ تم فتح {target}"
            })
        except Exception as e:
            return jsonify({
                "success": False,
                "message": f"فشل: {str(e)}"
            })
    
    # ============ إنشاء ملف نصي ============
    elif action == 'create_file':
        filename = data.get('filename', f'file_{datetime.now().strftime("%Y%m%d_%H%M%S")}.txt')
        content = data.get('content', '')
        filepath = os.path.expanduser(f'~/Desktop/{filename}')
        
        try:
            with open(filepath, 'w', encoding='utf-8') as f:
                f.write(content)
            return jsonify({
                "success": True,
                "message": f"✅ تم إنشاء {filename}"
            })
        except Exception as e:
            return jsonify({
                "success": False,
                "message": f"فشل: {str(e)}"
            })
    
    # ============ إنشاء ملف Excel ============
    elif action == 'create_excel':
        import openpyxl
        filename = data.get('filename', f'file_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx')
        columns = data.get('columns', [])
        filepath = os.path.expanduser(f'~/Desktop/{filename}')
        
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            
            # كتابة الأعمدة
            for i, col in enumerate(columns, 1):
                ws.cell(row=1, column=i, value=col)
            
            wb.save(filepath)
            return jsonify({
                "success": True,
                "message": f"✅ تم إنشاء {filename}"
            })
        except Exception as e:
            return jsonify({
                "success": False,
                "message": f"فشل: {str(e)}"
            })
    
    # ============ إنشاء ملف Word ============
    elif action == 'create_word':
        from docx import Document
        filename = data.get('filename', f'file_{datetime.now().strftime("%Y%m%d_%H%M%S")}.docx')
        content = data.get('content', '')
        filepath = os.path.expanduser(f'~/Desktop/{filename}')
        
        try:
            doc = Document()
            for line in content.split('\n'):
                doc.add_paragraph(line)
            doc.save(filepath)
            return jsonify({
                "success": True,
                "message": f"✅ تم إنشاء {filename}"
            })
        except Exception as e:
            return jsonify({
                "success": False,
                "message": f"فشل: {str(e)}"
            })
    
    # ============ بحث Google ============
    elif action == 'search':
        query = target or data.get('query', '')
        try:
            subprocess.Popen(['xdg-open', f'https://www.google.com/search?q={query}'])
            return jsonify({
                "success": True,
                "message": f"✅ بحث عن {query}"
            })
        except Exception as e:
            return jsonify({
                "success": False,
                "message": f"فشل: {str(e)}"
            })
    
    else:
        return jsonify({
            "success": False,
            "message": f"الأمر '{action}' غير معروف"
        })

@app.route('/health', methods=['GET'])
def health():
    return jsonify({"status": "ok"})

if __name__ == '__main__':
    print("🚀 Assistant Server started on http://localhost:5000")
    app.run(host='127.0.0.1', port=5000, debug=False)
